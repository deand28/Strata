"""GPU-free file grants, real parsers and request-selected native tools."""
import base64
import contextlib
import http.client
import io
import json
import os
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from serve.files import Files, FileHub, MAX_SOURCE, TOOLS, relative
from serve.frontend import ChatTemplate
from serve.harness import Harness, HarnessError
from serve.power import ModelPower
from serve.server import ByteTokenizer, Service, serve
from serve.test_mcp import ScriptedEngine, call_script

ROOT = Path(__file__).resolve().parents[1]


def pdf_bytes(blank=False):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    doc = PdfWriter()
    for number in range(2):
        page = doc.add_blank_page(width=600, height=800)
        if not blank:
            stream = DecodedStreamObject(); stream.set_data(f"BT /F1 12 Tf 20 700 Td (PDF COMET page {number + 1}) Tj ET".encode())
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            page[NameObject("/Contents")] = stream
    out = io.BytesIO(); doc.write(out); return out.getvalue()


class Local(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.base = Path(self.temp.name).resolve()
        self.docs = self.base / "docs"; self.docs.mkdir()
        (self.docs / "hello.txt").write_text("FOLDER MAPLE\n" + "x" * 50000, encoding="utf-8")
        (self.docs / "nested").mkdir(); (self.docs / "nested" / "notes.md").write_text("Nested content", encoding="utf-8")
        (self.base / "outside.txt").write_text("OUTSIDE SECRET", encoding="utf-8")
        self.files = Files(self.base / "settings")
        self.root = self.files.mutate({"action": "add", "path": str(self.docs)})["roots"][0]["id"]

    def tearDown(self): self.temp.cleanup()
    def read(self, path, **kwargs): return self.files.mutate({"action": "read", "root": self.root, "path": path, **kwargs})

    def test_explicit_persistent_grants_and_busy_remove(self):
        self.assertEqual(Files(self.base / "settings").snapshot(), self.files.snapshot())
        with self.files.guard(), self.assertRaises(HarnessError): self.files.mutate({"action": "remove", "root": self.root})
        self.files.mutate({"action": "remove", "root": self.root})
        with self.assertRaises(HarnessError): self.read("hello.txt")
        for value in (".", str(Path.home()), str(Path(self.docs.anchor)), "\\\\server\\share", "https://example.test/folder"):
            with self.subTest(path=value), self.assertRaises(HarnessError): self.files.mutate({"action": "add", "path": value})

    def test_escape_drive_ads_device_and_argument_types_rejected(self):
        for path in ("../outside.txt", "nested/../../outside.txt", "/outside.txt", "C:\\outside.txt", "\\\\server\\file", "hello.txt:secret", "nested//notes.md", "NUL.txt", "nested/notes.md.", "x\x00y"):
            with self.subTest(path=path), self.assertRaises(HarnessError): self.read(path)
        for value in (True, -1, 1.5, "1", 1000000001):
            with self.subTest(offset=value), self.assertRaises(HarnessError): self.read("hello.txt", offset=value)
        for value in (False, 0, 101, "10"):
            with self.subTest(count=value), self.assertRaises(HarnessError): self.read("hello.txt", count=value)

    def test_symlinks_and_junctions_never_traversed(self):
        link = self.docs / "outside-link"
        try: link.symlink_to(self.base, target_is_directory=True)
        except OSError:
            if os.name != "nt": raise
            result = subprocess.run(["cmd", "/c", "mklink", "/J", str(link), str(self.base)], capture_output=True)
            if result.returncode: self.skipTest("No symlink/junction capability")
        try:
            with self.assertRaises(HarnessError): self.read("outside-link/outside.txt")
            listing = self.files.mutate({"action": "list", "root": self.root})
            self.assertNotIn("outside-link", [r["name"] for r in listing["entries"]])
        finally:
            if link.is_symlink(): link.unlink()
            else: os.rmdir(link)

    def test_root_replaced_with_link_is_rejected(self):
        moved = self.base / "old"; self.docs.rename(moved)
        try: self.docs.symlink_to(moved, target_is_directory=True)
        except OSError: self.docs = moved; self.skipTest("No symlink privilege")
        try:
            with self.assertRaises(HarnessError): self.read("hello.txt")
        finally: self.docs.unlink()

    def test_directory_paging_name_search_and_lazy_read(self):
        value = self.files.mutate({"action": "list", "root": self.root, "count": 1})
        self.assertEqual(len(value["entries"]), 1); self.assertIsNotNone(value["next_offset"])
        other = self.files.mutate({"action": "list", "root": self.root, "count": 1, "offset": value["next_offset"]})
        self.assertNotEqual(value["entries"][0]["name"], other["entries"][0]["name"])
        found = self.files.mutate({"action": "search", "root": self.root, "query": "HELLO"})
        self.assertEqual([r["name"] for r in found["entries"]], ["hello.txt"])
        self.assertNotIn("Nested content", json.dumps(found))
        value = self.read("hello.txt"); self.assertIn("FOLDER MAPLE", value["text"]); self.assertEqual(value["next_offset"], 24000)
        self.assertLess(len(value["text"]), 24100)
        self.assertNotIn("FOLDER MAPLE", self.read("hello.txt", offset=value["next_offset"])["text"])

    def test_request_selection_enforced_at_native_call_boundary(self):
        h = Harness(self.base / "harness")
        try:
            req = {"messages": [{"role": "user", "content": "hello"}], "strata_files": {"roots": [self.root], "tools": ["files__read"]}}
            with h.request(req) as (prepared, hub), self.files.request(prepared, hub) as (prepared, chosen):
                self.assertTrue(prepared["strata_mcp"]); self.assertEqual(list(chosen.routes()), ["files__read"])
                self.assertIn("1 selected folders; enabled file tools: files__read", prepared["messages"][0]["content"])
                self.assertIn("FOLDER MAPLE", chosen.call("files__read", {"root": self.root, "path": "hello.txt"})["text"])
                self.assertFalse(chosen.call("files__list", {"root": self.root, "path": ""})["ok"])
                self.assertFalse(chosen.call("files__read", {"root": "r_000000000000", "path": "hello.txt"})["ok"])
                self.assertFalse(chosen.call("files__read", {"root": self.root, "path": "hello.txt", "count": True})["ok"])
                with self.assertRaises(HarnessError): self.files.mutate({"action": "remove", "root": self.root})
            with h.request({"messages": []}) as (prepared, hub), self.files.request(prepared, hub) as (prepared, chosen):
                self.assertFalse(prepared["strata_mcp"]); self.assertFalse(chosen.call("files__read", {})["ok"])
        finally: h.close()

    def test_utf8_and_utf16_pages_reconstruct_without_replacement_or_loss(self):
        for name, codec, marker in (("utf8.txt", "utf-8", "é😀"), ("utf16.txt", "utf-16", "🌳Birch")):
            source = marker * 16000
            (self.docs / name).write_bytes(source.encode(codec))
            pieces, offset = [], 0
            while offset is not None:
                value = self.read(name, offset=offset)
                pieces.append(value["text"].split("\n", 2)[2]); offset = value["next_offset"]
            self.assertEqual("".join(pieces), source)
            self.assertNotIn("\ufffd", "".join(pieces))
        with self.assertRaises(HarnessError): self.read("utf16.txt", offset=3)

    def test_capability_context_preserves_instructions_without_granting_access(self):
        h = Harness(self.base / "harness")
        try:
            for first in ({"role": "user", "content": "Can you read Word?"},
                          {"role": "system", "content": "Keep this instruction."},
                          {"role": "developer", "content": [{"type": "text", "text": "Keep this instruction."}]}):
                req = {"messages": [first]}
                before = json.dumps(req)
                with h.request(req) as (prepared, hub), self.files.request(prepared, hub) as (prepared, chosen):
                    text = json.dumps(prepared["messages"][0])
                    self.assertIn("Word DOCX", text)
                    self.assertIn("read-only", text)
                    self.assertIn("0 selected folders; enabled file tools: none", text)
                    self.assertFalse(prepared["strata_mcp"])
                    self.assertEqual(chosen.routes(), {})
                    self.assertFalse(chosen.call("files__read", {"root": self.root, "path": "hello.txt"})["ok"])
                    if first["role"] != "user": self.assertIn("Keep this instruction.", text)
                self.assertEqual(json.dumps(req), before)
            self.assertEqual(len(self.files.snapshot()["roots"]), 1, "capability context must not mutate grants")
        finally: h.close()

    def test_text_utf16_binary_unsupported_and_source_limits(self):
        text = self.files.extract("UTF16 BIRCH".encode("utf-16"), "notes.txt")
        self.assertIn("UTF16 BIRCH", text["text"])
        for data, name in ((b"\x00\x01binary", "bad.txt"), (b"legacy", "old.doc"), (b"x", "../notes.txt"), (b"x" * (MAX_SOURCE + 1), "large.pdf")):
            with self.subTest(name=name), self.assertRaises(HarnessError): self.files.extract(data, name)

    def test_unsupported_files_name_the_file_and_explain_video_and_metadata(self):
        for name, expected in (("lesson.mp4", ".mp4): Video reading/transcription is not available"),
                               ("old.doc", ".doc): This legacy Office format is unsupported"),
                               (".DS_Store", "macOS folder metadata"), ("archive.zip", ".zip): Unsupported file type")):
            with self.subTest(name=name):
                (self.docs / name).write_bytes(b"test")
                for read in (lambda: self.files.extract(b"test", name), lambda: self.read(name)):
                    with self.assertRaises(HarnessError) as error: read()
                    self.assertIn(name, str(error.exception)); self.assertIn(expected, str(error.exception))
        # A video above the document size limit still gets a precise unsupported-format diagnostic.
        with self.assertRaisesRegex(HarnessError, "lesson.mp4.*Video reading"):
            self.files.extract(b"x" * (MAX_SOURCE + 1), "lesson.mp4")

    def test_pdf_pages_and_image_only_are_explicit(self):
        data = pdf_bytes()
        first = self.files.extract(data, "report.pdf", count=1)
        self.assertIn("[Page 1]", first["text"]); self.assertIn("COMET page 1", first["text"]); self.assertNotIn("page 2", first["text"])
        self.assertEqual(first["next_offset"], 1)
        self.assertIn("[Page 2]", self.files.extract(data, "report.pdf", offset=1)["text"])
        self.assertIn("OCR is not included", self.files.extract(pdf_bytes(True), "scan.pdf")["warning"])

    def test_docx_tables_and_pptx_slides(self):
        from docx import Document
        d = Document(); d.add_paragraph("DOCX FALCON"); d.add_table(rows=1, cols=2).cell(0, 0).text = "Table Cedar"
        b = io.BytesIO(); d.save(b)
        result = self.files.extract(b.getvalue(), "proposal.docx", count=20)
        self.assertIn("DOCX FALCON", result["text"]); self.assertIn("Table Cedar", result["text"]); self.assertIn("[Table", result["text"])
        from pptx import Presentation
        p = Presentation()
        for word in ("PPTX SUNRISE", "Second slide"):
            s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = word
        b = io.BytesIO(); p.save(b)
        result = self.files.extract(b.getvalue(), "brief.pptx", count=1)
        self.assertIn("[Slide 1]", result["text"]); self.assertEqual(result["next_offset"], 1)
        self.assertNotIn("Second slide", result["text"])

    def test_xlsx_sheets_formulas_csv_tsv_and_rtf(self):
        import openpyxl
        d = openpyxl.Workbook(); d.active.title = "Budget"; d.active.append(["XLSX CEDAR", "=1+1"])
        d.create_sheet("Other").append(["Other sheet"]); b = io.BytesIO(); d.save(b); d.close()
        result = self.files.extract(b.getvalue(), "budget.xlsx", count=10)
        self.assertEqual(result["sheets"], ["Budget", "Other"]); self.assertIn("=1+1", result["text"])
        self.assertIn("[Sheet Budget · Row 1]", result["text"])
        self.assertIn("Other sheet", self.files.extract(b.getvalue(), "budget.xlsx", sheet="Other")["text"])
        for name, source in (("a.csv", b"name,value\nCSV ELM,25\nthird,3"), ("a.tsv", b"name\tvalue\nCSV ELM\t25\nthird\t3")):
            first = self.files.extract(source, name, count=1); self.assertEqual(first["next_offset"], 1)
            self.assertIn("CSV ELM", self.files.extract(source, name, offset=1)["text"])
        self.assertIn("RTF RIVER", self.files.extract(b"{\\rtf1\\ansi RTF RIVER\\par next}", "brief.rtf", count=100)["text"])

    def test_output_budget_preserves_whole_unit_and_rtf_cursors(self):
        from docx import Document
        d = Document()
        for i in range(4): d.add_paragraph(f"BLOCK-{i}-" + "x" * 10000)
        b = io.BytesIO(); d.save(b)
        first = self.files.extract(b.getvalue(), "long.docx", count=4)
        self.assertEqual(first["next_offset"], 2)
        self.assertIn("BLOCK-1-", first["text"]); self.assertNotIn("BLOCK-2-", first["text"])
        second = self.files.extract(b.getvalue(), "long.docx", offset=first["next_offset"], count=4)
        self.assertIn("BLOCK-2-", second["text"]); self.assertIn("BLOCK-3-", second["text"])
        source = "abcdefghij" * 6000
        data = ("{\\rtf1\\ansi " + source + "}").encode("ascii")
        parts, offset = [], 0
        while offset is not None:
            result = self.files.extract(data, "long.rtf", offset=offset)
            parts.append(result["text"].split("\n", 2)[2]); offset = result["next_offset"]
        self.assertEqual("".join(parts), source)

    def test_office_bomb_dtd_macros_malformed_and_worker_timeout(self):
        for member, data in (("word/document.xml", b"x" * 2_000_000), ("word/document.xml", b'<!DOCTYPE x [<!ENTITY a "secret">]><x>&a;</x>'), ("word/vbaProject.bin", b"macro")):
            b = io.BytesIO()
            with zipfile.ZipFile(b, "w", compression=zipfile.ZIP_DEFLATED) as z: z.writestr(member, data)
            with self.subTest(member=member), self.assertRaises(HarnessError): self.files.extract(b.getvalue(), "bad.docx")
        with self.assertRaises(HarnessError) as error: self.files.extract(b"bad private path C:/private.txt", "bad.pdf")
        self.assertNotIn("private", str(error.exception))
        with patch("serve.files.subprocess.run", side_effect=subprocess.TimeoutExpired(["fixed"], 15)), self.assertRaises(HarnessError) as error:
            self.files.extract(b"pdf", "slow.pdf")
        self.assertIn("15 seconds", str(error.exception))
        # Timeout releases parser ownership for the next request.
        self.assertTrue(self.files.parser_lock.acquire(False)); self.files.parser_lock.release()


class Http(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.base = Path(self.temp.name).resolve()
        self.docs = self.base / "docs"; self.docs.mkdir(); (self.docs / "proof.txt").write_text("FILE-TOOL-PROOF-73", encoding="utf-8")
        tok = ByteTokenizer(); self.engine = ScriptedEngine(tok, ["</think>\n\nDone."])
        self.svc = Service(self.engine, tok, ChatTemplate(ROOT / "serve/chat_template.jinja"))
        self.svc.harness = Harness(self.base / "state"); self.svc.files = Files(self.base / "state")
        self.httpd = serve(self.svc, port=0); self.port = self.httpd.server_address[1]
    def tearDown(self):
        self.httpd.shutdown(); self.httpd.server_close()
        if self.svc.harness: self.svc.harness.close()
        self.temp.cleanup()
    def call(self, path, body=None, headers=None):
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=25)
        try:
            c.request("GET" if body is None else "POST", path, body=None if body is None else (body if isinstance(body, bytes) else json.dumps(body).encode()), headers={"Content-Type": "application/json", **(headers or {})})
            r = c.getresponse(); return r.status, r.read().decode()
        finally: c.close()
    def test_management_and_attachment_work_off_without_wake(self):
        power = ModelPower(self.base / "state"); power.bind(self.svc); power.enabled = False; power.state = "off"
        with patch.object(self.svc, "load", side_effect=AssertionError("must not wake")):
            code, body = self.call("/files", {"action": "add", "path": str(self.docs)})
            self.assertEqual(code, 200, body); root = json.loads(body)["roots"][0]["id"]
            self.assertIn("proof.txt", self.call("/files", {"action": "list", "root": root})[1])
            code, body = self.call("/files/extract", {"name": "note.txt", "data": base64.b64encode(b"OFF ATTACHMENT").decode()})
            self.assertEqual(code, 200, body); self.assertIn("OFF ATTACHMENT", body)
            self.assertEqual(self.call("/v1/chat/completions", {"messages": []})[0], 503)
        self.assertFalse(power.enabled)
    def test_http_origin_auth_body_shape_limits_and_upstream_html(self):
        for path in ("/files", "/files/extract"):
            self.assertEqual(self.call(path, {}, {"Origin": "https://foreign.example"})[0], 403)
            self.assertEqual(self.call(path, {}, {"Content-Type": "text/plain"})[0], 415)
            self.assertEqual(self.call(path, b"not-json")[0], 400)
            self.assertEqual(self.call(path, [])[0], 400)
        self.assertEqual(self.call("/files/extract", {"name": "x.txt", "data": "@@@"})[0], 400)
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
        try:
            c.putrequest("POST", "/files/extract"); c.putheader("Content-Type", "application/json"); c.putheader("Content-Length", "99999999"); c.endheaders()
            r = c.getresponse(); self.assertEqual(r.status, 413); r.read()
        finally: c.close()
        self.svc.api_key = "test-key"; self.assertEqual(self.call("/files")[0], 401); self.svc.api_key = ""
        self.svc.harness = None; self.svc.files = None
        self.assertEqual(self.call("/files")[0], 404)
        self.assertEqual(self.call("/")[1], (ROOT / "serve/web/index.html").read_bytes().decode("utf-8"))
    def test_native_tool_real_loop_activity_context_and_unselected_absent(self):
        root = self.svc.files.mutate({"action": "add", "path": str(self.docs)})["roots"][0]["id"]
        self.engine = ScriptedEngine(self.svc.tok, [call_script("files__read", root=root, path="proof.txt"), "</think>\n\nRead complete."])
        self.svc.engine = self.engine
        code, raw = self.call("/v1/chat/completions", {"messages": [{"role": "user", "content": "Read proof.txt"}], "stream": True, "strata_files": {"roots": [root], "tools": ["files__read"]}})
        self.assertEqual(code, 200, raw)
        events = [json.loads(line[6:])["strata_mcp"] for line in raw.splitlines() if line.startswith("data: {") and "strata_mcp" in json.loads(line[6:])]
        self.assertEqual([e["event"] for e in events], ["start", "call", "result"])
        self.assertIn("FILE-TOOL-PROOF-73", events[-1]["text"])
        self.assertIn("File: proof.txt", self.engine.prompt_text(1))
        self.assertNotIn("files__search", self.engine.prompt_text(0))


if __name__ == "__main__": unittest.main()
