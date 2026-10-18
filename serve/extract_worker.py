"""One bounded, disposable document parser. Input is bytes, never a local path.

Invoked only by serve.files, with a fixed argv and no shell. Optional parsers are
imported after OS resource limits are established. Nothing is extracted to disk.
"""
from __future__ import annotations

import base64
import contextlib
import csv
import io
import json
import os
import re
import sys
import zipfile
from datetime import datetime
from pathlib import PurePath
from urllib.parse import urlsplit

MAX_SOURCE = 16 * 1024 * 1024
MAX_OUTPUT = 24000
LCPL_MAX = 128 * 1024
MEMORY = 512 * 1024 * 1024
_job = None


class DocumentError(ValueError):
    """Only deliberate parser diagnostics may cross the process boundary."""


def limits():
    global _job
    if os.name == "nt":
        import ctypes as c
        from ctypes import wintypes as w
        class IO(c.Structure):
            _fields_ = [(n, c.c_ulonglong) for n in ("ReadOperationCount", "WriteOperationCount", "OtherOperationCount", "ReadTransferCount", "WriteTransferCount", "OtherTransferCount")]
        class BASIC(c.Structure):
            _fields_ = [("PerProcessUserTimeLimit", c.c_longlong), ("PerJobUserTimeLimit", c.c_longlong), ("LimitFlags", w.DWORD), ("MinimumWorkingSetSize", c.c_size_t), ("MaximumWorkingSetSize", c.c_size_t), ("ActiveProcessLimit", w.DWORD), ("Affinity", c.c_size_t), ("PriorityClass", w.DWORD), ("SchedulingClass", w.DWORD)]
        class EXT(c.Structure):
            _fields_ = [("BasicLimitInformation", BASIC), ("IoInfo", IO), ("ProcessMemoryLimit", c.c_size_t), ("JobMemoryLimit", c.c_size_t), ("PeakProcessMemoryUsed", c.c_size_t), ("PeakJobMemoryUsed", c.c_size_t)]
        k = c.WinDLL("kernel32", use_last_error=True)
        k.CreateJobObjectW.argtypes = [c.c_void_p, w.LPCWSTR]; k.CreateJobObjectW.restype = w.HANDLE
        k.SetInformationJobObject.argtypes = [w.HANDLE, c.c_int, c.c_void_p, w.DWORD]; k.SetInformationJobObject.restype = w.BOOL
        k.GetCurrentProcess.restype = w.HANDLE
        k.AssignProcessToJobObject.argtypes = [w.HANDLE, w.HANDLE]; k.AssignProcessToJobObject.restype = w.BOOL
        _job = k.CreateJobObjectW(None, None)
        info = EXT(); info.BasicLimitInformation.LimitFlags = 0x100 | 0x2 | 0x8  # memory, CPU, active-process count
        info.BasicLimitInformation.PerProcessUserTimeLimit = 10 * 10000000
        info.BasicLimitInformation.ActiveProcessLimit = 1
        info.ProcessMemoryLimit = MEMORY
        if not _job or not k.SetInformationJobObject(_job, 9, c.byref(info), c.sizeof(info)) or not k.AssignProcessToJobObject(_job, k.GetCurrentProcess()):
            raise DocumentError("Document parser resource limits are unavailable on this system")
    else:
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (MEMORY, MEMORY))
        resource.setrlimit(resource.RLIMIT_CPU, (10, 11))
        resource.setrlimit(resource.RLIMIT_FSIZE, (0, 0))
        resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))


def office_preflight(data):
    from defusedxml.ElementTree import fromstring
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        entries = z.infolist()
        if len(entries) > 2000 or sum(i.file_size for i in entries) > 32 * 1024 * 1024:
            raise DocumentError("Office document exceeds the expanded-size limit")
        for entry in entries:
            if entry.flag_bits & 1 or entry.file_size > 8 * 1024 * 1024 or entry.file_size > max(1024 * 1024, entry.compress_size * 200):
                raise DocumentError("Office document contains an oversized or encrypted part")
            name = entry.filename.lower()
            if "vbaproject" in name or "/embeddings/" in name:
                raise DocumentError("Office documents containing macros or embedded binaries are unsupported")
            if name.endswith((".xml", ".rels")):
                fromstring(z.read(entry), forbid_dtd=True, forbid_entities=True, forbid_external=True)


def clipped(value, limit=MAX_OUTPUT):
    return str(value)[:limit]


def lcp_license(data, name, offset):
    """Describe a lending license without disclosing account/key/link fields."""
    if len(data) > LCPL_MAX:
        raise DocumentError("LCP license exceeds the 128 KiB limit")
    if offset:
        raise DocumentError("LCP licenses have one metadata range; use offset 0")
    try:
        license = json.loads(data.decode("utf-8-sig"))
    except (UnicodeError, ValueError):
        raise DocumentError("Invalid LCP license JSON") from None
    if not isinstance(license, dict) or not isinstance(license.get("encryption"), dict) or not isinstance(license.get("links"), list):
        raise DocumentError("Invalid LCP license structure")
    publications = [link for link in license["links"][:100] if isinstance(link, dict) and
                    (link.get("rel") == "publication" or isinstance(link.get("rel"), list) and "publication" in link["rel"])]
    if not publications:
        raise DocumentError("LCP license has no publication link")
    kinds = {"application/pdf+lcp": "PDF", "application/epub+zip": "EPUB", "application/audiobook+lcp": "audiobook"}
    formats = sorted({kinds.get(link.get("type"), "protected publication") for link in publications if isinstance(link.get("type"), str)})
    provider = license.get("provider")
    try:
        parsed = urlsplit(provider) if isinstance(provider, str) else None
        host = parsed.hostname if parsed and parsed.scheme in ("http", "https") else None
        if not host or len(host) > 255 or not all(ch.isascii() and (ch.isalnum() or ch in ".-") for ch in host):
            host = None
    except ValueError:
        host = None
    rights = license.get("rights") if isinstance(license.get("rights"), dict) else {}
    lines = ["[Readium LCP borrowing license; publication text is not in this file]",
             "Protected publication: " + ", ".join(formats or ["format unspecified"])]
    if host:
        lines.append("Provider: " + host)
    for label, value in (("Issued", license.get("issued")), ("Loan starts", rights.get("start")), ("Loan ends", rights.get("end"))):
        if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d{1,9})?(?:Z|[+-]\d{2}:\d{2})", value):
            try:
                datetime.fromisoformat(value.replace("Z", "+00:00"))
            except ValueError:
                continue
            lines.append(label + ": " + value)
    lines.append("This license does not contain the book pages. Open the .lcpl in a Readium LCP-compatible reader such as Thorium Reader to access the loan. Strata has not downloaded or decrypted the publication or checked the license status/signature.")
    return {"name": name, "text": f"File: {name}\n" + "\n".join(lines), "unit": "license",
            "offset": 0, "next_offset": None, "total": 1, "truncated": False,
            "warning": "License metadata only; no book text was read. Open the .lcpl in an LCP-compatible reader."}


def extract(data, name, offset=0, count=4, sheet=None):
    ext = PurePath(name).suffix.lower()
    lines, next_offset, warning = [], None, ""
    if ext == ".lcpl":
        return lcp_license(data, name, offset)
    if ext in (".docx", ".xlsx", ".pptx"):
        office_preflight(data)
    if ext == ".pdf":
        from pypdf import PdfReader
        doc = PdfReader(io.BytesIO(data), strict=True)
        if doc.is_encrypted:
            raise DocumentError("Encrypted PDFs are unsupported; provide a decrypted copy")
        total = len(doc.pages)
        if total > 2000:
            raise DocumentError("PDF exceeds the 2,000-page limit; split it into smaller documents")
        end = min(total, offset + min(count, 8))
        for i in range(offset, end):
            text = doc.pages[i].extract_text() or ""
            lines.append(f"[Page {i + 1}]\n{clipped(text)}")
        unit = "page"; next_offset = end if end < total else None
        if not any(x.strip() for x in [s.split("\n", 1)[-1] for s in lines]):
            warning = "No extractable text on these pages. This may be a scanned/image-only PDF; OCR is not included."
    elif ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        # XML body order keeps tables beside the paragraphs that introduce them.
        from docx.oxml.ns import qn
        blocks = list(doc.element.body)
        total = len(blocks); end = min(total, offset + min(count, 100))
        for i in range(offset, end):
            block = blocks[i]
            words = [n.text or "" for n in block.iter(qn("w:t"))]
            kind = "Table" if block.tag == qn("w:tbl") else "Paragraph"
            lines.append(f"[{kind} {i + 1}]\n{clipped(' | '.join(words))}")
        unit = "block"; next_offset = end if end < total else None
    elif ext == ".xlsx":
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            names = book.sheetnames
            if sheet is not None and sheet not in names:
                raise DocumentError("Sheet not found; choose one of the returned sheet names")
            chosen = sheet if sheet is not None else names[0]
            ws = book[chosen]
            total = min(ws.max_row or 0, 100000)
            if offset > 100000: raise DocumentError("Row offset exceeds the 100,000-row limit")
            end = min(total, offset + min(count, 100))
            for row in ws.iter_rows(min_row=offset + 1, max_row=max(offset + 1, end), max_col=min(ws.max_column or 1, 100), values_only=True):
                if offset >= end: break
                lines.append(f"[Sheet {chosen} · Row {offset + len(lines) + 1}] " + " | ".join(clipped(v if v is not None else "", 2000) for v in row))
            unit = "row"; next_offset = end if end < total else None
            warning = "Formulas are shown as text and are never calculated. At most 100 columns and 100,000 rows are accessible."
        finally:
            book.close()
        extra = {"sheets": names, "sheet": chosen}
    elif ext == ".pptx":
        from pptx import Presentation
        doc = Presentation(io.BytesIO(data)); total = len(doc.slides)
        end = min(total, offset + min(count, 10))
        for i in range(offset, end):
            parts = []
            for shape in doc.slides[i].shapes:
                if shape.has_text_frame: parts.append(clipped(shape.text))
                if shape.has_table:
                    parts.extend(" | ".join(clipped(c.text, 2000) for c in row.cells) for row in shape.table.rows)
            lines.append(f"[Slide {i + 1}]\n{clipped(chr(10).join(parts))}")
        unit = "slide"; next_offset = end if end < total else None
    elif ext in (".csv", ".tsv"):
        csv.field_size_limit(32768)
        source = decode(data)
        rows = csv.reader(io.StringIO(source), delimiter="\t" if ext == ".tsv" else ",")
        total = 0; end = offset + min(count, 100); more = False
        for i, row in enumerate(rows):
            total = i + 1
            if i >= end: more = True; break
            if i >= offset: lines.append(f"[Row {i + 1}] " + " | ".join(clipped(v, 2000) for v in row[:100]))
        unit = "row"; next_offset = end if more else None
    elif ext == ".rtf":
        if len(data) > 4 * 1024 * 1024: raise DocumentError("RTF exceeds the 4 MiB limit")
        from striprtf.striprtf import rtf_to_text
        text = rtf_to_text(data.decode("latin-1"), errors="replace")
        total = len(text); end = min(total, offset + MAX_OUTPUT - 64)
        lines.append(f"[Characters {offset}–{end}]\n{text[offset:end]}")
        unit = "character"; next_offset = end if end < total else None
    else:
        raise DocumentError("Unsupported document type")
    # Stop at a whole-unit boundary so the cursor never jumps past unreturned
    # pages/rows/blocks/slides. A single oversized unit is explicitly clipped.
    kept, used, partial = [], 0, False
    for line in lines:
        separator = 2 if kept else 0
        if used + separator + len(line) > MAX_OUTPUT:
            if not kept:
                kept.append(line[:MAX_OUTPUT]); partial = True
            break
        kept.append(line); used += separator + len(line)
    if len(kept) < len(lines):
        next_offset = offset + len(kept)
        warning = (warning + " Range stopped at the output limit; continue with next_offset.").strip()
    if partial:
        warning = (warning + " Text within this single page/block/row/slide was clipped to 24,000 characters; split this part to read it in full.").strip()
    content = "\n\n".join(kept)
    result = {"name": name, "text": f"File: {name}\n" + content, "unit": unit,
              "offset": offset, "next_offset": next_offset, "total": total,
              "truncated": partial or next_offset is not None, "warning": warning}
    if not content.strip() or not any(part.split("]", 1)[-1].strip() for part in lines): result["warning"] = warning or "No extractable text in this range."
    if ext == ".xlsx": result.update(extra)
    return result


def decode(data):
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        text = data.decode("utf-16")
    else:
        text = data.decode("utf-8-sig")
    if "\x00" in text or sum(ord(c) < 32 and c not in "\n\r\t\f" for c in text) > max(2, len(text) // 100):
        raise DocumentError("This is binary data, not a supported text document")
    return text


def main():
    original = sys.stdout
    try:
        limits()
        request = json.loads(sys.stdin.buffer.read(24 * 1024 * 1024 + 1))
        data = base64.b64decode(request["data"], validate=True)
        if len(data) > MAX_SOURCE: raise DocumentError("Document exceeds the 16 MiB limit")
        with open(os.devnull, "w") as sink, contextlib.redirect_stdout(sink), contextlib.redirect_stderr(sink):
            result = extract(data, request["name"], request["offset"], request["count"], request.get("sheet"))
        payload = {"result": result}
    except ImportError:
        payload = {"error": "Document readers are not installed. Install requirements-files.txt in the server's Python environment."}
    except DocumentError as e:
        payload = {"error": str(e)}
    except BaseException:
        payload = {"error": "Document could not be read. It may be malformed, encrypted, or exceed parser resource limits."}
    output = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    if len(output) > 180000:
        output = json.dumps({"error": "Document metadata exceeds the output limit; split it into smaller files"}).encode("utf-8")
    original.buffer.write(output); original.buffer.flush()


if __name__ == "__main__": main()
